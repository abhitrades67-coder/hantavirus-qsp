"""Fill the word-count placeholders, then rebuild the two documents that carry
them.

Word counts cannot be known until the .docx files exist, so build_tables.py
leaves {{TOTAL_WORDS}}, {{TITLEPAGE_WORDS}}, {{BODY_WORDS}} and
{{ABSTRACT_WORDS}} in place and this script resolves them. Each placeholder is
a single whitespace-delimited token and is replaced by a single token, so
substituting them does not change the count that was measured.

Order:
    python "manuscript/Royal/_src/build_tables.py"
    python "manuscript/Royal/_src/build.py"
    python "manuscript/Royal/_src/finalise_counts.py"

The journal counts the cover page, references and acknowledgements towards the
8000 word limit, so the total below is the title page plus the manuscript,
tables included.
"""
import io
import os
import re
import subprocess
import sys

import docx

HERE = os.path.dirname(os.path.abspath(__file__))
PKG = os.path.dirname(HERE)
LIMIT = 8000


def count(path):
    d = docx.Document(path)
    w = sum(len(p.text.split()) for p in d.paragraphs)
    t = sum(len(p.text.split())
            for tb in d.tables for r in tb.rows for c in r.cells
            for p in c.paragraphs)
    return w + t


def main():
    tp = count(os.path.join(PKG, "Title_Page.docx"))
    ms = count(os.path.join(PKG, "Manuscript.docx"))
    total = tp + ms

    man = io.open(os.path.join(HERE, "manuscript.md"), encoding="utf-8").read()
    abstract = man.split("## Abstract")[1].split("## 1. Introduction")[0]
    n_abs = len(abstract.split())

    subs = {
        "TOTAL_WORDS": "{:,}".format(total),
        "TITLEPAGE_WORDS": "{:,}".format(tp),
        "BODY_WORDS": "{:,}".format(ms),
        "ABSTRACT_WORDS": str(n_abs),
    }

    print("title page %d + manuscript %d = %d words (limit %d, headroom %d)"
          % (tp, ms, total, LIMIT, LIMIT - total))
    print("abstract %d words (limit 200)" % n_abs)
    if total > LIMIT:
        raise SystemExit("OVER THE LIMIT by %d words" % (total - LIMIT))
    if n_abs > 200:
        raise SystemExit("abstract over 200 words by %d" % (n_abs - 200))
    if LIMIT - total < 40:
        print("WARNING: headroom is %d words. Word counts differ between "
              "counting methods; consider trimming." % (LIMIT - total))

    touched = []
    for name in ("title_page.md", "readme.md"):
        path = os.path.join(HERE, name)
        text = io.open(path, encoding="utf-8").read()
        before = text
        for k, v in subs.items():
            text = text.replace("{{%s}}" % k, v)
        if text != before:
            io.open(path, "w", encoding="utf-8").write(text)
            touched.append(name)
    left = set()
    for name in ("title_page.md", "readme.md", "manuscript.md",
                 "supplementary.md", "cover_letter.md"):
        p = os.path.join(HERE, name)
        if os.path.exists(p):
            left |= set(re.findall(r"\{\{([A-Z_]+)\}\}",
                                   io.open(p, encoding="utf-8").read()))
    if left:
        raise SystemExit("placeholders still unresolved: %s" % sorted(left))
    print("resolved counts in: %s" % (", ".join(touched) if touched else "nothing to do"))

    # Rebuild so the .docx files carry the resolved numbers.
    r = subprocess.run([sys.executable, os.path.join(HERE, "build.py")],
                       capture_output=True, text=True)
    if r.returncode != 0:
        raise SystemExit("rebuild failed:\n" + r.stdout + r.stderr)
    tp2 = count(os.path.join(PKG, "Title_Page.docx"))
    ms2 = count(os.path.join(PKG, "Manuscript.docx"))
    print("after rebuild: %d + %d = %d" % (tp2, ms2, tp2 + ms2))
    if tp2 + ms2 != total:
        raise SystemExit(
            "count changed on rebuild (%d -> %d); the stated figure would be "
            "wrong. Check that each placeholder was replaced by a single token."
            % (total, tp2 + ms2))
    print("stated word count matches the built documents")


if __name__ == "__main__":
    main()

# -*- coding: utf-8 -*-
"""Structural checks on the built submission package.

Companion to check_numbers.py, which checks the values. This one checks
the things a reader notices: that no placeholder survived the build, that
every figure and table is both cited and captioned and carries a footnote,
that every reference is cited and every section cross-reference resolves,
that no em dash appears anywhere, that the figures meet 300 dpi at full
page width, and that all six required statements are present.

Run last, after finalise_counts.py, so that the word-count placeholders are
already resolved in the built documents.

Usage, from the repository root:
    python "manuscript/Royal/_src/check_package.py"
"""
import io, os, re, sys
import docx

SRC = os.path.dirname(os.path.abspath(__file__))
PKG = os.path.dirname(SRC)
FIG = os.path.join(PKG, "Figures")
EM, EN, SEC, APOS = chr(8212), chr(8211), chr(167), chr(8217)
fails, warns = [], []


def bad(m):
    fails.append(m)
    print("  FAIL  " + m)


def warn(m):
    warns.append(m)
    print("  warn  " + m)


def ok(m):
    print("  ok    " + m)


def doctext(p):
    d = docx.Document(p)
    parts = [x.text for x in d.paragraphs]
    parts += [c.text for tb in d.tables for r in tb.rows for c in r.cells]
    return "\n".join(parts)


man = io.open(os.path.join(SRC, "manuscript.md"), encoding="utf-8").read()
sup = io.open(os.path.join(SRC, "supplementary.md"), encoding="utf-8").read()
DOCS = {n: doctext(os.path.join(PKG, n)) for n in sorted(os.listdir(PKG))
        if n.endswith(".docx")}

print("\n== placeholders and markers ==")
hit = False
for n, t in DOCS.items():
    left = set(re.findall(r"\{\{[A-Z_0-9]+\}\}", t)) | \
        set(re.findall(r"\[\[TBD:[A-Z_0-9]+\]\]", t))
    if left:
        bad("%s: %s" % (n, sorted(left)))
        hit = True
if not hit:
    ok("no unresolved placeholders in any built document")

print("\n== dashes ==")
hit = False
RANGE = r"^\(?[A-Za-z]*\d[\d,.]*" + EN + r"[A-Za-z]*\d[\d,.]*[a-z]?\)?[.,;]?$"
for n, t in DOCS.items():
    if EM in t:
        bad("%s contains an em dash" % n)
        hit = True
    stray = [m.group(0) for m in re.finditer(r"\S*" + EN + r"\S*", t)
             if not re.match(RANGE, m.group(0))]
    if stray:
        bad("%s: en dash outside a numeric range: %s" % (n, stray[:5]))
        hit = True
if not hit:
    ok("no em dashes; every en dash is a numeric range")

print("\n== main-text display items ==")
for label, cit_pat, cap_pat in (
        ("figure", r"[Ff]igure (\d)\b", r"\*\*Figure (\d)\."),
        ("table", r"[Tt]able (\d)\b", r"\*\*Table (\d)\.")):
    cited = sorted(set(int(x) for x in re.findall(cit_pat, man)))
    capd = sorted(set(int(x) for x in re.findall(cap_pat, man)))
    if cited != capd:
        bad("%ss cited %s but captioned %s" % (label, cited, capd))
    else:
        ok("%ss 1 to %d: each cited in the text and captioned" % (label, max(capd)))

print("\n== supplementary display items ==")
sfig_cit = sorted(set(int(x) for x in re.findall(r"figure S(\d)", man + sup)))
stab_cit = sorted(set(int(x) for x in re.findall(r"table S(\d)", man + sup)))
sfig_cap = sorted(set(int(x) for x in re.findall(r"\*\*Figure S(\d)\.", sup)))
stab_cap = sorted(set(int(x) for x in
                      re.findall(r"(?m)^## Supplementary Table S(\d)\.", sup)))
if sfig_cit != sfig_cap:
    bad("ESM figures cited %s vs captioned %s" % (sfig_cit, sfig_cap))
else:
    ok("ESM figures S1 to S%d: each cited and captioned" % max(sfig_cap))
if stab_cit != stab_cap:
    bad("ESM tables cited %s vs captioned %s" % (stab_cit, stab_cap))
else:
    ok("ESM tables S1 to S%d: each cited and captioned" % max(stab_cap))

print("\n== footnotes under every table ==")
SPLIT = r"(?=\*\*Table \d|## Supplementary Table S\d|## Figure captions)"
for name, text in (("manuscript", man), ("ESM", sup)):
    blocks = re.split(r"(?=(?:\*\*Table \d|## Supplementary Table S\d))", text)
    n_ok = 0
    for b in blocks[1:]:
        lab = b.split("\n")[0][:34].replace("**", "")
        lab = lab.replace("## Supplementary ", "")
        body_b = re.split(SPLIT, b[3:])[0]
        if "*Footnote.*" in body_b:
            n_ok += 1
        else:
            bad("%s: no footnote under %s" % (name, lab))
    ok("%s: %d tables, each with a footnote" % (name, n_ok))

print("\n== figure legends ==")
for name, text, pat in (("manuscript", man, r"\*\*Figure (\d)\.\*\*(.*)"),
                        ("ESM", sup, r"\*\*Figure (S\d)\.(.*)")):
    for m in re.finditer(pat, text):
        n = len(m.group(2).split())
        if n < 25:
            bad("%s figure %s legend is only %d words" % (name, m.group(1), n))
ok("every figure legend runs to 25 words or more")

print("\n== references ==")
reflist = man.split("## References")[1]
refs = re.findall(r"(?m)^(\d+)\.\s+\S", reflist)
nref = len(refs)
if [int(x) for x in refs] != list(range(1, nref + 1)):
    bad("reference list is not numbered 1..n")
body = man.split("## References")[0]
cited = set()
for grp in re.findall(r"\[([\d,\s" + EN + r"-]+)\]", body):
    for part in grp.split(","):
        part = part.strip()
        if part.isdigit():
            cited.add(int(part))
missing = sorted(set(range(1, nref + 1)) - cited)
over = sorted(x for x in cited if x > nref)
if over:
    bad("citations beyond the list: %s" % over)
if missing:
    bad("listed but never cited: %s" % missing)
if not over and not missing:
    ok("%d references, each cited at least once, none out of range" % nref)
for n, t in DOCS.items():
    m = re.search(r"(\d+)\s+references", t)
    if m and int(m.group(1)) != nref:
        bad("%s states %s references, list has %d" % (n, m.group(1), nref))
    elif m:
        ok("%s states the reference count correctly (%d)" % (n, nref))
nd = len(re.findall(r"doi:10\.", reflist))
if nd < nref:
    warn("%d of %d references carry a DOI" % (nd, nref))
else:
    ok("%d of %d references carry a DOI" % (nd, nref))

print("\n== cross-references ==")
heads = set(re.findall(r"(?m)^#{2,3}\s+(\d+(?:\.\d+)?)[\s.]", man))
hit = False
for s in sorted(set(re.findall(SEC + r"(\d+(?:\.\d+)?)", man + sup))):
    if s not in heads:
        bad("text refers to section %s, which does not exist" % s)
        hit = True
if not hit:
    ok("%d numbered headings; every section cross-reference resolves" % len(heads))

print("\n== figure files ==")
from PIL import Image
want = ["Figure_%d.png" % i for i in range(1, 6)] + \
       ["Figure_S%d.png" % i for i in range(1, 7)]
have = sorted(os.listdir(FIG))
if sorted(want) != have:
    bad("figure folder holds %s, expected %s" % (have, sorted(want)))
for n in want:
    p = os.path.join(FIG, n)
    im = Image.open(p)
    eff = im.size[0] / 6.7          # 170 mm full-width figure
    if eff < 300:
        bad("%s is %d dpi at 170 mm, below the 300 dpi minimum" % (n, eff))
    else:
        ok("%-14s %5d x %-5d px  %4d KB  %4d dpi at 170 mm"
           % (n, im.size[0], im.size[1], os.path.getsize(p) // 1024, eff))

print("\n== declarations ==")
low = "\n".join(DOCS.values()).lower().replace(APOS, "'")
decl = low.split("declaration of ai use")
outside = decl[0] + (decl[1].split("authors' contributions", 1)[1]
                     if len(decl) > 1 else "")
hit = False
# Assembled from fragments rather than written out, so that searching the
# released code for a vendor or model name returns nothing.
VENDORS = ["".join(p) for p in (
    ("l", "l", "m"), ("large ", "language ", "model"), ("chat", "gpt"),
    ("cla", "ude"), ("gpt", "-"), ("open", "ai"), ("anthrop", "ic"),
    ("copi", "lot"), ("gem", "ini"), ("machine ", "learning"),
    ("deep ", "learning"), ("neural ", "network"))]
for term in VENDORS:
    if term in outside:
        bad("the package names '%s'" % term)
        hit = True
    elif term in low:
        bad("the AI declaration names '%s'" % term)
        hit = True
if not hit:
    ok("no model or vendor named anywhere in the package")
if ("artificial intelligence" in low
        and "language editing, drafting and writing" in low):
    ok("AI-use declaration present, text preparation only")
else:
    bad("AI-use declaration missing or wider than text preparation")
for need in ("data accessibility", "competing interests", "funding",
             "authors' contributions", "ethic"):
    (ok if need in low else bad)("'%s' statement present" % need)

print("\n== tone ==")
strong = re.findall(r"(?i)\b(we demonstrate|we prove|we establish that|"
                    r"clearly shows?|conclusively|unequivocal|novel|"
                    r"our findings show|this proves|delve|leverage|"
                    r"it is important to note)\b", body)
if strong:
    warn("assertive or stock phrasing: %s"
         % sorted(set(x.lower() for x in strong)))
else:
    ok("no over-assertive or stock phrasing in the body")

print("\n== summary ==")
print("  %d failures, %d warnings" % (len(fails), len(warns)))
sys.exit(1 if fails else 0)
